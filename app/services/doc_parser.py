"""
文档解析服务：从上传的简历文档（PDF/DOCX/TXT/MD）提取文本，
再用 LLM 解析为结构化简历数据。
"""
import io
import json
import logging
from typing import Dict, Any

from .llm_service import llm_service

logger = logging.getLogger(__name__)


class DocParserService:
    """文档解析 + LLM 结构化提取"""

    SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md"}

    # ------------------------------------------------------------------
    # 文本提取
    # ------------------------------------------------------------------
    def extract_text(self, filename: str, content: bytes) -> str:
        """根据文件类型提取纯文本"""
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in self.SUPPORTED_EXT:
            raise ValueError(
                f"不支持的文件格式: {ext}，请上传 PDF/DOCX/TXT/MD 文件"
            )
        if ext == ".pdf":
            return self._extract_pdf(content)
        elif ext == ".docx":
            return self._extract_docx(content)
        else:
            return self._extract_text(content)

    def _extract_pdf(self, content: bytes) -> str:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise RuntimeError("未安装 PyPDF2，请联系管理员安装")
        reader = PdfReader(io.BytesIO(content))
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)
        return "\n".join(pages).strip()

    def _extract_docx(self, content: bytes) -> str:
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("未安装 python-docx，请联系管理员安装")
        doc = Document(io.BytesIO(content))
        # 1) 段落文本
        chunks = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        # 2) 表格文本（教育/经历常以表格形式存在）
        for ti, tbl in enumerate(doc.tables):
            rows_text = []
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    rows_text.append(" | ".join(cells))
            if rows_text:
                chunks.append(f"\n[表格{ti+1}]\n" + "\n".join(rows_text))
        return "\n".join(chunks).strip()

    def _extract_text(self, content: bytes) -> str:
        for encoding in ("utf-8", "gbk", "gb2312", "big5"):
            try:
                return content.decode(encoding).strip()
            except (UnicodeDecodeError, LookupError):
                continue
        return content.decode("utf-8", errors="replace").strip()

    # ------------------------------------------------------------------
    # LLM 结构化解析
    # ------------------------------------------------------------------
    def parse_resume(self, text: str) -> Dict[str, Any]:
        """调用 LLM 把简历文本解析为结构化 JSON 数据"""
        system_prompt = (
            "你是专业的中文简历解析器。任务：从用户给出的简历纯文本中提取结构化信息，"
            "以纯JSON输出，禁止输出任何解释、说明、markdown代码块。\n"
            "\n"
            "【输出字段定义】（不存在的填 null 或空数组 []，不要省略字段）\n"
            "name: 姓名（必填，中文姓名通常2-4字，位于简历顶部或个人信息的'姓名:'标签后）\n"
            "title: 求职意向/职位（如'AIAgent开发工程师'）\n"
            "email, phone, location, github, blog: 联系方式\n"
            "summary: 自我评价/个人简介\n"
            "educations: 教育经历数组，每项 {school, degree, major, start_date, end_date, description}\n"
            "  - school: 学校全称（如'周口师范学院''北京大学'），识别关键词: XX学院/XX大学/XX学校/XX研究院\n"
            "  - degree: 学位（本科/硕士/博士/大专/学士）\n"
            "  - major: 专业\n"
            "internships: 实习经历数组，每项 {company, position, start_date, end_date, location, highlights[]}\n"
            "projects: 项目经验数组，每项 {name, role, start_date, end_date, description, highlights[], tech_stack[], link}\n"
            "skills: 技能数组，每项 {category, items[]}\n"
            "\n"
            "【硬性规则】\n"
            "1. 必须从原文中提取真实姓名作为name，不要写'未命名'或null（除非真的没有）\n"
            "2. 教育经历必须提取school字段，识别关键词: 'XX学院''XX大学''XX学校'\n"
            "3. 字段名全部用英文，日期保持原文格式\n"
            "4. 只输出纯JSON对象，{开头}结尾，无任何额外字符\n"
            "\n"
            "【输出JSON示例】\n"
            "{\n"
            '  "name": "张峻宇",\n'
            '  "title": "AIAgent开发工程师",\n'
            '  "email": "zhang@example.com",\n'
            '  "phone": "13800138000",\n'
            '  "location": "河南郑州",\n'
            '  "github": "github.com/xxx",\n'
            '  "blog": null,\n'
            '  "summary": "2年AI Agent开发经验...",\n'
            '  "educations": [{"school":"周口师范学院","degree":"本科","major":"计算机科学与技术","start_date":"2020.09","end_date":"2024.06","description":null}],\n'
            '  "internships": [{"company":"XX科技有限公司","position":"AI开发实习生","start_date":"2023.07","end_date":"2023.12","location":"北京","highlights":["负责LLM应用开发"]}],\n'
            '  "projects": [{"name":"智能客服Agent","role":"主要开发","start_date":"2024.01","end_date":"2024.06","description":"基于LLM的智能客服系统","highlights":["使用LangChain编排"],"tech_stack":["Python","LangChain","FastAPI"],"link":null}],\n'
            '  "skills": [{"category":"编程语言","items":["Python","JavaScript"]},{"category":"AI框架","items":["LangChain","OpenAI"]}]\n'
            "}\n"
        )
        user_prompt = (
            "请解析以下简历文本并提取结构化JSON。\n"
            "要求：1)必须提取真实姓名 2)必须识别学校全称 3)只输出纯JSON：\n"
            "\n"
            "===== 简历文本开始 =====\n"
            f"{text[:8000]}\n"
            "===== 简历文本结束 =====\n"
            "\n"
            "请输出纯JSON："
        )

        messages = llm_service.build_messages(system_prompt, user_prompt)
        content, tokens = llm_service.chat_completion(
            messages, temperature=0.0, max_tokens=6000
        )

        cleaned = self._clean_json(content)
        try:
            data = json.loads(cleaned)
        except json.JSONDecodeError as e:
            logger.error(f"LLM返回JSON解析失败: {e}\n原始内容: {content[:500]}")
            raise RuntimeError("AI解析简历失败，返回的内容不是合法JSON，请重试或手动录入")

        # 清洗 null 值：LLM 可能返回 null 但 Pydantic 必填字段不接受
        if not data.get("name"):
            data["name"] = "未命名"
        data.setdefault("educations", [])
        data.setdefault("internships", [])
        data.setdefault("projects", [])
        data.setdefault("skills", [])
        # 清洗子表中必填字段的 null
        for edu in data["educations"]:
            if not edu.get("school"):
                edu["school"] = "未知学校"
        for intern in data["internships"]:
            if not intern.get("company"):
                intern["company"] = "未知公司"
        for proj in data["projects"]:
            if not proj.get("name"):
                proj["name"] = "未知项目"
        for sk in data["skills"]:
            if not sk.get("items"):
                sk["items"] = []
        logger.info(f"[DocParser] LLM解析完成 name={data.get('name')} "
                     f"edu={len(data['educations'])} int={len(data['internships'])} "
                     f"proj={len(data['projects'])} skill={len(data['skills'])}")
        return data

    def _clean_json(self, text: str) -> str:
        """去掉 LLM 可能包裹的 markdown 代码块标记"""
        text = text.strip()
        if text.startswith("```"):
            lines = text.split("\n")
            if lines[0].startswith("```"):
                lines = lines[1:]
            if lines and lines[-1].strip() == "```":
                lines = lines[:-1]
            text = "\n".join(lines)
        return text.strip()


# 全局单例
doc_parser = DocParserService()
